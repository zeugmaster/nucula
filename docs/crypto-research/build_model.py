#!/usr/bin/env python3
"""Build an explicitly hypothetical model alongside recorded research evidence."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


HERE = Path(__file__).resolve().parent
BASELINE = 1.886479
PHASES = [
    ("Incoming BLS verification", 1.886479),
    ("Spend-info reconstruction", 0.166180),
    ("Signing and witness JSON (includes duplicate signing)", 0.832368),
    ("Output blinding", 0.400463),
    ("Output unblinding", 0.404532),
    ("Returned-proof BLS verification", 1.886468),
]


def style_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="343D46")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_view.showGridLines = False


def csv_sheet(book, title, paths):
    sheet = book.create_sheet(title)
    first = True
    for path in paths:
        with path.open(newline="") as handle:
            rows = csv.reader(handle)
            header = next(rows)
            if first:
                sheet.append(["source_file"] + header)
                first = False
            for row in rows:
                values = [int(x) if x.isdigit() else x for x in row]
                sheet.append([path.name] + values)
    style_sheet(sheet, [42, 39, 26, 23, 24])


def main():
    book = Workbook()
    book.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
    intro = book.active
    intro.title = "Read me"
    intro.append(["Item", "Meaning"])
    for row in [
        ("Status", "Research model, 2026-09-13. No new ESP32-C3 timing measurements."),
        ("Historical observations", "September 1 recorded synthetic Nutroot receive; old proposal, omitted ECDH/output key generation, duplicate signing."),
        ("Model", "T = baseline × [(1-f) + f × (1-r)/s]. f is affected field-kernel time share; r removes calls; s accelerates remaining calls."),
        ("Assumptions", "Blue cells on Assumptions are editable. f is unknown until profiled. Default r=20% is an illustrative approximation, not an integrated firmware measurement."),
        ("Sensitivity", "Nine independent scenarios, with their own explicit f and s, reference r and baseline from Assumptions. Results are conditional calculations."),
        ("Operation counts", "32-bit-limb host probe; excludes software inverse internals, additions, hashing, CPU/control and memory traffic."),
        ("Stack frames", "Cross-compiler estimates for individual functions, not complete task high-water marks."),
        ("Nutroot counts", "Four well-formed fixtures of the older local core. Experimental fast paths are not production patches."),
        ("Units", "Time in seconds, memory in bytes, count columns are integer call counts."),
        ("Scope", "See ../crypto-optimization-esp32c3.md, README.md, and provenance.json for analysis, reproduction and source fingerprints."),
    ]:
        intro.append(row)
    style_sheet(intro, [30, 112])
    for i in range(2, intro.max_row + 1):
        intro.row_dimensions[i].height = 43

    history = book.create_sheet("Historical observations")
    history.append(["Phase", "Recorded seconds", "Qualification"])
    for label, seconds in PHASES:
        history.append([label, seconds, "Historical synthetic fixture; see Read me"])
    history.append(["Sum", "=SUM(B2:B7)", "Not a current-proposal end-to-end measurement"])
    history.append(["BLS share", "=(B2+B7)/B8", "Fraction of this historical sum"])
    history.append(["Only BLS phases 2x faster", "=B8-(B2+B7)/2", "Amdahl illustration; all other phases held fixed"])
    style_sheet(history, [64, 24, 62])
    for row in range(2, 11):
        history.cell(row, 2).number_format = "0.000000"
    history["B9"].number_format = "0.0%"

    assumptions = book.create_sheet("Assumptions")
    assumptions.append(["Parameter", "Value", "Meaning"])
    for row in [
        ("baseline_seconds", "='Historical observations'!B2", "Historical incoming verification only"),
        ("f_kernel_time_fraction", 0.70, "UNKNOWN. Illustrative affected-kernel share; must be measured."),
        ("r_kernel_calls_removed", 0.20, "Illustrative fraction removed; not a wall-clock saving."),
        ("s_kernel_speedup", 1.50, "HYPOTHETICAL speedup of each remaining affected-kernel call."),
        ("selected_result_seconds", "=B2*((1-B3)+B3*(1-B4)/B5)", "Conditional result, not a hardware prediction."),
        ("selected_overall_speedup", "=B2/B6", "Baseline divided by selected result."),
        ("count_accounting_example", (22150 + 4494 + 4724) / 159730, "Cold G2 checks + chunk change + forced MSM; separate probe components, not integrated firmware."),
    ]:
        assumptions.append(row)
    style_sheet(assumptions, [34, 25, 100])
    for name in ["B3", "B4", "B5"]:
        assumptions[name].font = Font(color="1565C0", bold=True)
        assumptions[name].fill = PatternFill("solid", fgColor="EAF3FC")
    for name in ["B3", "B4", "B8"]:
        assumptions[name].number_format = "0.0%"
    for name in ["B2", "B6"]:
        assumptions[name].number_format = "0.000000"
    for name in ["B5", "B7"]:
        assumptions[name].number_format = '0.00"x"'

    scenarios = book.create_sheet("Sensitivity")
    headers = ["assumed_f", "assumed_s", "assumed_r", "baseline_seconds",
               "conditional_seconds", "overall_speedup", "time_reduction_fraction"]
    scenarios.append(headers)
    numeric = []
    for f in [0.50, 0.70, 0.85]:
        for s in [1.0, 1.5, 2.0]:
            row = scenarios.max_row + 1
            scenarios.append([f, s, "=Assumptions!$B$4", "=Assumptions!$B$2",
                f"=D{row}*((1-A{row})+A{row}*(1-C{row})/B{row})",
                f"=D{row}/E{row}", f"=1-E{row}/D{row}"])
            seconds = BASELINE * ((1-f) + f * (1-0.20) / s)
            numeric.append([f, s, 0.20, BASELINE, seconds, BASELINE/seconds, 1-seconds/BASELINE])
    style_sheet(scenarios, [19, 19, 19, 24, 25, 22, 26])
    for row in range(2, scenarios.max_row + 1):
        for col in [1, 3, 7]:
            scenarios.cell(row, col).number_format = "0.0%"
        for col in [4, 5]:
            scenarios.cell(row, col).number_format = "0.000000"
        for col in [2, 6]:
            scenarios.cell(row, col).number_format = '0.00"x"'
    with (HERE / "sensitivity.csv").open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(numeric)

    csv_sheet(book, "Field operation counts", sorted(HERE.glob("operation-counts-*.csv")))
    csv_sheet(book, "Stack frames", [HERE / "stack-frames.csv"])
    csv_sheet(book, "Nutroot call counts", sorted(HERE.glob("nutroot-*.csv")))
    sources = book.create_sheet("Sources")
    sources.append(["Source", "Location / qualification"])
    sources.append(["Historical timings", "../bench-nutroot-esp32c3.md"])
    sources.append(["Local source provenance", "provenance.json"])
    sources.append(["Interpretation and full citations", "../crypto-optimization-esp32c3.md"])
    sources.append(["BLS proposal status", "https://github.com/cashubtc/nuts/pull/371"])
    sources.append(["Nutroot proposal status", "https://github.com/cashubtc/nuts/pull/421"])
    for row in sources.iter_rows(min_row=2):
        if str(row[1].value).startswith("https:"):
            row[1].hyperlink = row[1].value
            row[1].font = Font(color="1565C0", underline="single")
    style_sheet(sources, [38, 100])
    book.save(HERE / "performance-model.xlsx")
    print("Wrote performance-model.xlsx and sensitivity.csv; hypothetical results are explicitly labeled.")


if __name__ == "__main__":
    main()
